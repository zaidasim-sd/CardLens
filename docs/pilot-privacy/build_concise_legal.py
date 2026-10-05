from pathlib import Path
import unicodedata
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'output/documents/CardSnap Product Privacy Terms and Aventure Agreement.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
for name,size in [('Normal',9),('Title',12),('Heading 2',10.5),('List Bullet',9)]:
    s=doc.styles[name]
    s.font.name='Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    s.font.bold=name in ['Title','Heading 2']
    s.paragraph_format.space_after=Pt(5)
    s.paragraph_format.space_before=Pt(7 if name=='Heading 2' else 0)
    s.paragraph_format.line_spacing=1.08
    fonts=s.element.get_or_add_rPr().rFonts
    for k in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+k),'Aptos')
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: fonts.attrib.pop(qn('w:'+k),None)
for node in list(doc.styles.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
# Use an explicit round bullet rather than a theme dependent list marker.
numbering=doc.part.numbering_part.element
abstract=OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'),'70')
lvl=OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'),'0')
for tag,val in [('start','1'),('numFmt','bullet'),('lvlText','•'),('lvlJc','left')]:
    el=OxmlElement('w:'+tag); el.set(qn('w:val'),val); lvl.append(el)
pp=OxmlElement('w:pPr'); ind=OxmlElement('w:ind'); ind.set(qn('w:left'),'240'); ind.set(qn('w:hanging'),'180'); pp.append(ind); lvl.append(pp)
rp=OxmlElement('w:rPr'); ff=OxmlElement('w:rFonts'); ff.set(qn('w:ascii'),'Aptos'); ff.set(qn('w:hAnsi'),'Aptos'); rp.append(ff); lvl.append(rp)
abstract.append(lvl); numbering.append(abstract)
num=OxmlElement('w:num'); num.set(qn('w:numId'),'70'); ai=OxmlElement('w:abstractNumId'); ai.set(qn('w:val'),'70'); num.append(ai); numbering.append(num)

def p(t,style=None):
    assert ';' not in t and not any(unicodedata.category(c)=='Pd' for c in t)
    return doc.add_paragraph(t,style)
def title(t): p(t,'Title')
def h(t): p(t,'Heading 2')
def b(t):
    para=p(t,'List Bullet')
    props=para._p.get_or_add_pPr(); np=OxmlElement('w:numPr')
    for tag,value in [('ilvl','0'),('numId','70')]:
        el=OxmlElement('w:'+tag); el.set(qn('w:val'),value); np.append(el)
    props.append(np)
def page(t): doc.add_page_break(); title(t)

title('CardSnap Privacy Notice')
b('Draft for review dated 29 September 2026. This notice covers CardSnap as a standard Vision71 product. Complete the marked details and confirm the deployed setup before publication or real personal data use.')
h('Responsibility and purpose')
b('Vision71 [legal name, address and privacy email] provides CardSnap. Each customer organisation controls the business contact data it collects. Vision71 processes that data on the customer’s documented instructions under the applicable agreement.')
b('CardSnap extracts business card text, supports human review and checks local duplicates. Each customer must establish its lawful basis and provide its own contact privacy notice.')
b('A card, camera permission or verified status does not establish marketing permission. The application does not send marketing messages or make significant automated decisions about people.')
h('Information and recipients')
b('Information includes card images, filenames, extracted text, suggested and reviewed contact fields, meeting notes, identifiers, status and dates.')
b('Vercel hosts the service. Google Cloud Vision is the intended primary OCR provider. OCR.space, operated by a9t9 software GmbH, can receive the same image on fallback. Both may receive a card.')
b('Manual entry and the prepared demo bypass OCR. Constant Contact is not connected and receives no information through the current application.')
b('Confirm production provider settings, contracting entities and any support, logging or backup recipients before publication.')
h('Storage and retention')
b('There is no server contact database in the inspected app. The backend handles images in memory without an application file or storage bucket write.')
b('Submitting a reviewed contact saves its image, extracted text and details in persistent browser IndexedDB. Closing the browser does not delete them. The code does not use localStorage for these records.')
b('Saved records have no automatic expiry or deletion after extraction or review. Set and test retention periods for browser records, originals, logs, support copies and backups before live use [periods and owners].')
b('Google describes online image processing in memory with temporary metadata logs. OCR.space states that images are deleted after processing and API IP logs are kept for one month. Host logs require separate confirmation.')
h('Location security and rights')
b('The website is https://cardsnapbyv71.vercel.app/. Vercel Functions default to iad1 in Washington, D.C., USA. The actual project region is unverified. Static files may be served globally. Confirm OCR locations and any required transfer safeguards.')
b('Provider calls use HTTPS and backend API keys. The current app has no login or enforced roles. People using the same browser profile can access saved records. Confirm the final access controls before live use.')
b('Camera permission is optional. No analytics or advertising integration was found in the app code. Confirm hosting additions and cookies before publication.')
b('Depending on applicable law, people may request access, correction, deletion, restriction or portability, object to processing and withdraw consent where used. Contact [privacy email]. Complaints may be made to [relevant regulator and link].')
b('Send privacy requests to the organisation that collected your card or to [Vision71 privacy email] for routing. Report suspected exposure to [incident email and telephone]. Access and support arrangements follow the customer’s service agreement. Update this notice when the setup changes. Effective date: [complete].')

page('CardSnap Terms of Use')
b('Draft for review. These terms apply to customers and authorised users of CardSnap, a standard Vision71 product. Provider: Vision71 [legal name, address and contact]. Effective date and acceptance method: [complete].')
h('Permitted use')
b('Use CardSnap only for authorised business card capture and review. Until the service is approved for live personal data, use synthetic data only, including manual entries.')
b('Submit only information you may lawfully collect and share. Do not submit sensitive information, identity documents, payment credentials or unrelated private notes.')
b('Check extracted details and duplicate warnings before using them. Do not misuse the service, bypass access controls or send unlawful messages.')
h('Current service')
b('The app sends images to configured OCR providers and saves submitted records in the current browser. It has no server contact database, application login or enforced roles.')
b('There is no separate manager approval. A verified status reflects an operator action. Hiding queue details does not anonymise the records.')
b('Constant Contact is not connected. The current export button has no working action. Future integrations require written approval.')
b('Browser records persist until removed. Use approved devices and follow the agreed retention process. Report lost devices and accidental disclosures promptly.')
h('Service commitments')
b('Vision71 will perform agreed work with reasonable care and skill. OCR may be inaccurate. Provider failures, quotas and network problems can interrupt scanning.')
b('Fees, service periods, support hours, usage limits and any availability or recovery commitments must be agreed in writing. The current service is not a reliable backup.')
b('Before business use, confirm a hosting arrangement that permits commercial activity. Vercel’s Hobby plan is restricted to personal, noncommercial use.')
h('Data and product rights')
b('The customer retains its rights in client data. Vision71 may process it only for the agreed service and lawful instructions. No sale, unrelated advertising or model training is authorised.')
b('Vision71 retains its rights in CardSnap, subject to any separate written agreement. Customers may use the product within their agreed service scope. Source code delivery, ownership transfer and independent hosting rights require a separate agreement. External software licences continue to apply.')
b('Possessing a card or accepting these terms does not establish marketing permission. Any contact transfer requires an approved purpose, recipient and field mapping.')
h('Suspension termination and disputes')
b('Vision71 may suspend affected processing when reasonably necessary to contain an incident, prevent unlawful use or address material breach. Notify the customer promptly where lawful.')
b('Ending the service does not remove agreed handover, return or deletion duties. Removing online access does not erase copies on devices.')
b('Nothing excludes liability that cannot lawfully be excluded or removes statutory rights. Complete agreed liability limits, exceptions, governing law and courts before acceptance.')
b('A signed customer agreement takes priority over conflicting public terms. Its Data Processing Schedule governs personal data matters. Material changes require the agreed notice and approval process.')

page('Aventure Data and Service Agreement with Data Processing Schedule')
b('Draft for client and legal review. Proposed obligations are not evidence of existing controls. No live data is authorised until the launch conditions are met.')
h('Parties and pilot')
b('This agreement applies only to Aventure Aviation, referred to as Aventure, and Vision71. Full legal names, addresses and authorised contacts: [complete]. Pilot dates, users, devices, volume, fees, support and governing law: [complete].')
b('Scope: business card capture, OCR and human contact review. Constant Contact and other integrations are excluded unless separately agreed.')
b('Before launch, approve the facts sheet, deployed version, regions, providers, contracts, access controls, retention, deletion, incident process and public notices. Test controls with synthetic data and obtain Aventure’s written approval.')
h('Data Processing Schedule')
b('Aventure acts as controller and establishes lawful collection, notices, marketing permissions and retention. Vision71 acts as processor on documented instructions, including transfer instructions.')
b('Purpose: Aventure business contact management. People: cardholders, business representatives and people named in permitted notes. Data: images, filenames, OCR text, contact fields, notes, identifiers, status and dates.')
b('Operations: receive images, transmit for OCR, parse, display, correct, check duplicates, save locally and perform approved return or deletion. Duration: pilot term plus only the agreed exit period.')
b('No intentional collection of sensitive information or children’s data is authorised. Vision71 must promptly flag unlawful instructions. Notify Aventure before legally required processing outside instructions unless prohibited.')
b('Vision71 must limit access to authorised personnel bound to confidentiality. No sale, independent marketing or unrelated model training is permitted.')
h('Providers and protection')
b('Vercel, Google Cloud Vision and OCR.space require a completed schedule of entities, accounts, plans, countries and contracts. Obtain Aventure’s written approval before use or replacement. Apply equivalent protection duties to subprocessors. Vision71 remains responsible for delegated processor duties.')
b('Confirm actual processing locations and any required international transfer mechanism and safeguards. The Vercel default region does not establish Google or OCR.space residency.')
b('Before live use, test authentication covering the app and OCR endpoint, appropriate permissions, privileged account protection, HTTPS, secret restrictions, request limits and device security.')
b('Agree minimal activity records, log access and expiry, and backup and recovery arrangements. The current app has no login, enforced roles, user audit trail or working application backup process.')
h('Retention and transfers')
b('There is no server contact database. Submitted images and details persist in browser IndexedDB without automatic expiry. Do not promise temporary storage or automatic deletion until implemented and tested.')
b('Complete retention periods, deletion triggers and owners for browser records, original photos, exports, logs, support copies and backups. Deletion must include images, original OCR and reviewed details. Masking and archiving are not deletion.')
b('A future Constant Contact transfer needs an agreed method, account owner, minimum field mapping, permission evidence, suppression handling, error checks and revocation process. Test it with synthetic data first.')

page('Aventure Incident Response and Product Handover')
h('Incidents and individual rights')
b('Vision71 must notify Aventure without undue delay after becoming aware of a breach. Proposed maximum: 24 hours, subject to agreement. Fill and test the monitored email and urgent telephone contacts before launch.')
b('Provide known scope, likely effects, containment and the next update. Assist investigation and required notifications. Aventure decides regulator and individual notices unless law places a separate duty on Vision71. Statutory deadlines still apply.')
b('Promptly forward individual requests and assist access, correction, deletion and other applicable rights. Assist security assessments, impact assessments and regulator consultation. Provide compliance information and allow proportionate audits.')
h('Complete product handover')
b('Deliver source code, agreed history, build and deployment instructions, configuration inventory, dependency licences and operating guidance through a secure channel.')
b('Agree rights allowing Aventure to run, copy, modify and maintain the product, including through a replacement provider. Record any payment conditions or ownership assignment expressly.')
b('Transfer hosting, cloud, OCR, repository and deployment control to Aventure. Where an account cannot be transferred, migrate to an Aventure owned account and test it before retiring the old service.')
b('Aventure must control ownership, billing, recovery methods and administrator invitations. Aventure must issue replacement passwords and keys that Vision71 cannot access.')
b('Remove Vision71 users, inherited administrator roles, service accounts, sessions, tokens and deployment keys. Remove Vision71 monitoring, log forwarding, webhooks, remote support and deployment connections. Retire old deployments that could process client data.')
b('Verify with synthetic data that the service uses Aventure’s accounts, sends no information to Vision71 destinations and rejects removed identities and credentials. Aventure must approve the final access roster and handover in writing.')
b('After verified handover, Vision71 must retain no administrative or operational access to Aventure’s scanning activity, card images or extracted details. Approved hosting and OCR providers will still process information to operate the service.')
b('Any later Vision71 support access requires Aventure’s explicit written approval, a defined purpose, minimum permissions and an expiry time, followed by removal.')
h('Return deletion and agreement')
b('Aventure chooses return followed by deletion, or deletion without return, of data held by Vision71. Proposed completion: within 30 calendar days of termination or the agreed instruction. Revoke routine access at handover.')
b('Delete Vision71 copies and provide a signed completion record. Identify any legally required retention. Backup exceptions need a fixed expiry and restricted use. Aventure manages retention on its own devices. Confidentiality and required deletion duties continue after termination.')
b('Changes to purposes, providers, countries, storage or access require written approval and updated notices. Complete termination rights, notice periods, transition costs, liability terms and governing law before signature. Mandatory legal rights remain unaffected.')
b('The Data Processing Schedule prevails on personal data matters. Mandatory law and applicable transfer clauses take priority where required.')
b('Vision71 authorised name, signature and date: [complete]. Aventure authorised name, signature and date: [complete]. Technical approval and legal review: [complete].')

footer=sec.footer.paragraphs[0]; footer.alignment=2
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='CardSnap Privacy Terms and Aventure Agreement'
doc.core_properties.author='Vision71'
doc.save(OUT)
print(OUT)
