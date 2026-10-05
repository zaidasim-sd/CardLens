from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"D:\CardLens\docs\CardSnap_Ali_Bhai_Setup_Request_and_Architecture.docx"
BLUE = "315C85"
TEXT = "1F2937"
PALE = "F4F7FA"
BORDER = "D9D9D9"


def set_font(run, size=9, bold=False, color=TEXT, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, top=75, start=100, bottom=75, end=100):
    props = cell._tc.get_or_add_tcPr()
    container = props.first_child_found_in("w:tcMar")
    if container is None:
        container = OxmlElement("w:tcMar")
        props.append(container)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = container.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            container.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_borders(table):
    props = table._tbl.tblPr
    container = props.first_child_found_in("w:tblBorders")
    if container is None:
        container = OxmlElement("w:tblBorders")
        props.append(container)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = container.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            container.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), BORDER)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def add_title(doc, title, subtitle):
    paragraph = doc.add_paragraph(style="Title")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(title)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    set_font(paragraph.add_run(subtitle), 9, True, BLUE)


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.add_run(text)
    return paragraph


def add_bullet(doc, text, size=8.8):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.02
    set_font(paragraph.add_run(text), size)
    return paragraph


def add_number(doc, text, size=8.8):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.02
    set_font(paragraph.add_run(text), size)
    return paragraph


def add_table(doc, headers, widths, rows, body_size=7.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.text = ""
        shade(cell, BLUE)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(paragraph.add_run(header), 8, True, "FFFFFF")
    repeat_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            cell.width = Inches(widths[column_index])
            cell.text = ""
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cell, PALE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 0.98
            set_font(paragraph.add_run(value), body_size, column_index == 0, "000000" if column_index == 0 else TEXT)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.48)
section.left_margin = Inches(0.6)
section.right_margin = Inches(0.6)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.04

for style_name, size in (("Title", 18.5), ("Heading 1", 11.3)):
    style = doc.styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(6 if style_name == "Heading 1" else 0)
    style.paragraph_format.space_after = Pt(3)

# Page 1
add_title(doc, "CardSnap Setup Request for Ali Bhai", "Required cloud setup and secure credential handoff for Version 1")
paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_after = Pt(6)
set_font(paragraph.add_run("Request. "), 9.4, True, "000000")
set_font(paragraph.add_run("Please create or approve the accounts and resources below, then place the required credentials directly in the protected Vercel and GitHub settings. This will allow Haroon and Hassan to complete the live Google Vision and Google Sheets tests. Credentials must not be sent through WhatsApp, chat, browser code or the repository."), 9.4)

add_heading(doc, "Selected services and Version 1 decisions")
decisions = [
    ("1  Login and authentication", "CardSnap named accounts using Node scrypt password hashes, server stored sessions, Secure SameSite Strict cookies, CSRF protection, rate limits, lockout and immediate session revocation. No paid identity provider is required."),
    ("2  Database and image storage", "MongoDB Atlas Free. Contact fields, OCR text and temporary images are encrypted with AES 256 GCM. Images are stored separately in the cardImages collection and never kept in browser storage."),
    ("3  Proposed region", "Mumbai is the first choice when available for both application compute and Atlas. Keep application and database services in the same geographic area. Ali and Aventure must approve the final region and data residency position."),
    ("4  Image retention", "24 hours by default. Vision71 administration may set 0 to 168 hours. Zero means no image is saved. The hourly audited sweep is primary and MongoDB TTL is a two hour safety net."),
    ("5  Backup and restore", "Run an encrypted backup daily through GitHub Actions and retain the encrypted artifact for 30 days. Store it outside Atlas. Restore only into an empty separate database and compare restored collection counts. The automated restore test has passed."),
    ("6  Google Sheets connection", "Use the Vision71 CardSnap service account through the backend. Share only the selected Sheet as Editor with cardsnap@cardsnap-510210.iam.gserviceaccount.com. Store its JSON private key only in protected environment settings."),
    ("7  Aventure API key", "Aventure does not need to provide a Google API key. Vision71 supplies the controlled service account. A separate Vision71 Google Vision key is required for OCR and must be restricted to Cloud Vision only."),
    ("8  Version 1 workflow", "CardSnap capture, Aventure owned Google Sheet, Hala reviews with a status dropdown, Approved rows are exported as CSV, then Aventure or Vision71 manually imports the approved contacts into the agreed Constant Contact list. There is no automatic Constant Contact connection."),
]
add_table(doc, ("Decision", "Selected approach"), (1.85, 5.15), decisions, 7.35)

add_heading(doc, "Current status")
paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_after = Pt(0)
set_font(paragraph.add_run("Complete: "), 8.5, True, "000000")
set_font(paragraph.add_run("private staging deployment from haroon, named test accounts, roles, encrypted storage, audit controls, retention, backup and restore, protected OCR logic, Sheet mapping and status sync, Approved CSV export, 112 passing automated tests and a passing production build."), 8.5)

# Page 2
doc.add_page_break()
add_title(doc, "Setup Steps and Internal Test Plan", "Actions required before Ali Bhai’s staging review")
add_heading(doc, "Setup steps for Ali Bhai")
steps = [
    "MongoDB Atlas. Create or confirm one Free cluster, select the approved region, create the cardsnap_vision71_staging database and create a database user restricted to that database. Configure the network allowlist for Vercel after reviewing the lack of fixed Vercel outbound addresses. Place MONGODB_URI directly in Vercel and the approved local environment.",
    "Google Cloud. Use project cardsnap-510210, enable the Google Sheets API and Cloud Vision API, and confirm the service account cardsnap@cardsnap-510210.iam.gserviceaccount.com is enabled.",
    "Service account credential. Create a JSON key for the CardSnap service account. Place client_email as GOOGLE_SERVICE_ACCOUNT_EMAIL and private_key as GOOGLE_SERVICE_ACCOUNT_PRIVATE_KEY directly in Vercel. Do not send the JSON file through chat or commit it.",
    "Google Sheet. For internal testing, share the Vision71 test Sheet with the service account as Editor. For live use, Aventure creates or copies the approved template into its own Drive and shares only that Sheet. Set the approved Sheet URL or identifier in the protected environment.",
    "Google Vision. Revoke the key previously exposed in chat. Create a replacement restricted to Cloud Vision only and place it directly in Vercel as GOOGLE_VISION_API_KEY. CardSnap sends it in the X Goog Api Key header and never in a URL.",
    "Backup and sweep. Add APP_BASE_URL, CRON_SECRET, MONGODB_URI and BACKUP_KEY to GitHub Actions. Confirm BACKUP_KEY differs from ENCRYPTION_KEY.",
    "Final integration. Haroon and Hassan run the Sheet configuration command, submit fake records, confirm Hala’s status updates return to CardSnap, test Approved CSV export, run mobile checks and record the short demonstration video."
]
for item in steps:
    add_number(doc, item, 8.35)

add_heading(doc, "Secure values required")
secure_values = [
    "MONGODB_URI for the restricted CardSnap database user",
    "GOOGLE_SERVICE_ACCOUNT_EMAIL and GOOGLE_SERVICE_ACCOUNT_PRIVATE_KEY",
    "A replacement GOOGLE_VISION_API_KEY restricted to Cloud Vision only",
    "The approved Google Sheet URL or identifier and tab identifier",
    "CRON_SECRET, ENCRYPTION_KEY and a different BACKUP_KEY"
]
for value in secure_values:
    add_bullet(doc, value, 8.35)

paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_after = Pt(5)
set_font(paragraph.add_run("Handoff method. "), 8.5, True, "000000")
set_font(paragraph.add_run("Ali Bhai should enter these values directly in the Vercel and GitHub secret settings or use an approved one time secret manager. Haroon and Hassan need access to use the configured values, not copies stored in messages or files."), 8.5)

add_heading(doc, "Internal testing dates owners and evidence")
schedule = [
    ("1 October 2026", "Ali Bhai with Haroon and Hassan", "Cloud accounts, restricted identities and environment names visible in provider settings. No secret values in screenshots."),
    ("2 October 2026", "Haroon and Hassan", "Live Google Vision and Sheet write tests using fake contacts, captured request results and audit references."),
    ("3 to 4 October 2026", "Zaid and Ibrahim as Capturers, Hala as Reviewer", "Capture, duplicate, missing field, correction, approval, rejection, failed Sheet retry and removed user results in the test report."),
    ("5 October 2026", "Haroon and Hassan", "Mobile capture and review evidence, Approved CSV, backup restore evidence and short demonstration video."),
    ("6 October 2026", "Ali Bhai", "Private staging review against the security checklist and recorded approval or required changes.")
]
add_table(doc, ("Date", "Owner", "Required evidence"), (1.25, 1.75, 4.0), schedule, 7.35)

paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_before = Pt(5)
paragraph.paragraph_format.space_after = Pt(0)
set_font(paragraph.add_run("Proposed staging review: 6 October 2026. "), 8.6, True, "000000")
set_font(paragraph.add_run("This date depends on the protected credentials and Sheet access being configured by 1 October 2026."), 8.6)

# Page 3
doc.add_page_break()
add_title(doc, "CardSnap Version 1 One Page Recommendation", "Decision owner Ali Bhai  |  Internal testing followed by the first Aventure exhibition")
paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_after = Pt(6)
set_font(paragraph.add_run("Status: "), 8.8, True, "000000")
set_font(paragraph.add_run("Recommendation only. No production resource should be created until written approval."), 8.8)

architecture = [
    ("Login and authentication", "Use the CardSnap named account system already implemented. It uses Node scrypt password hashes, server stored sessions, Secure and SameSite Strict cookies, CSRF protection, sign in rate limits, account lockout, eight hour maximum sessions, thirty minute idle expiry and immediate session revocation when an account is removed. Vision71 creates the initial named accounts. Hala reviews in the Sheet and does not need a CardSnap login for Version 1."),
    ("Database", "Use a dedicated MongoDB Atlas Free database for internal testing. Keep every tenant separate and encrypt contact fields, OCR text and images in the application with AES 256 GCM. Use a database user restricted to the CardSnap database. Atlas Free has a 512 MB limit, no automatic backups and no database audit, so usage monitoring, application audit records and external encrypted backups are required."),
    ("Image storage", "Store compressed images temporarily in the encrypted MongoDB cardImages collection for Version 1. Do not create Google Cloud Storage or another production storage resource until Ali approves it. Images must not remain in browser storage."),
    ("Cloud provider and region", "Keep the application in a separate Vercel staging project. For the live pilot, propose a region near the operating team, with Mumbai as the first choice if it is available for both application compute and Atlas. Keep the application and database in the same geographic area where possible. Ali and Aventure must approve the final region and data residency position before live data is used."),
    ("Retention and deletion", "Use a default card image retention period of 24 hours, configurable from 0 to 168 hours by Vision71 administration. A value of 0 prevents image storage. Run the audited deletion sweep every hour and use the MongoDB TTL index as a two hour safety net. Contact records remain until an approved deletion request or the agreed end of pilot deletion."),
    ("Backup method", "Run the Node backup script daily through GitHub Actions. Export cards, users without password hashes, lists and settings. Exclude images, sessions, login attempts, rate limits and tokens. Compress and encrypt each backup with a separate 32 byte backup key, retain the encrypted artifact for 30 days and perform a documented restore test after material storage changes."),
    ("Expected monthly cost", "Expected internal test cost is 0 USD while usage remains inside approved free allowances. This assumes MongoDB Atlas Free, existing GitHub Actions allowance and Google Vision usage within its approved allowance. Vercel Hobby is restricted to non commercial use, so it must not be treated as an approved commercial production plan. Google Cloud billing requirements and the live hosting cost must be confirmed before the pilot. Stop before enabling any resource that requires an unapproved payment method or charge."),
    ("Who can access the data", "Capturers can submit contacts and view or correct only their own records. Hala can view and update the protected Aventure owned Sheet. Vision71 Administrators can manage named users, retention, deletion and Approved CSV export. Vision71 Support can view aggregate counts only, cannot browse contact details or images and expires after 24 hours. The backend service account can access only the approved Sheet. Constant Contact receives data only through the approved manual CSV import in Version 1.")
]
add_table(doc, ("Area", "Recommendation"), (1.48, 5.52), architecture, 7.18)
add_heading(doc, "Approval points")
paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_after = Pt(0)
set_font(paragraph.add_run("Ali and Aventure must approve the final cloud region, live hosting plan, Atlas network access, image retention period, Google service account access, Aventure owned Sheet and any expected charge before real Aventure data is used."), 8.2)

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer_paragraph.add_run("CardSnap by Vision71  |  30 September 2026"), 7.2, False, "6B7280")

properties = doc.core_properties
properties.title = "CardSnap Setup Request and Production Architecture"
properties.subject = "Setup request and decision note for Ali Bhai"
properties.author = "Vision71"
properties.keywords = "CardSnap, Google Sheets, MongoDB, production architecture"
doc.save(OUT)
print(OUT)
