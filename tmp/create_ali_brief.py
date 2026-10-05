from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r'D:\CardLens\docs\CardSnap_Ali_Bhai_Google_Sheets_and_Production_Architecture.docx'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_table_borders(table, color='D9D9D9', size='5'):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = f'w:{edge}'
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:color'), color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_font(run, name='Aptos', size=9, bold=False, color='1F2937'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(text), size=8.4)
    return p

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.48)
section.bottom_margin = Inches(0.48)
section.left_margin = Inches(0.58)
section.right_margin = Inches(0.58)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
normal.font.size = Pt(8.7)
normal.font.color.rgb = RGBColor(31, 41, 55)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.02

for style_name, size in [('Title', 19), ('Heading 1', 11.5)]:
    s = styles[style_name]
    s.font.name = 'Aptos Display'
    s._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
    s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.space_before = Pt(6 if style_name == 'Heading 1' else 0)
    s.paragraph_format.space_after = Pt(3)

p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.space_after = Pt(2)
p.add_run('CardSnap Google Sheets Access and Production Architecture')

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
r = p.add_run('One page decision note for Ali Bhai  |  Version 1 internal testing and first exhibition')
set_font(r, size=9, bold=True, color='315C85')

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
r = p.add_run('Decision requested. ')
set_font(r, size=9.2, bold=True, color='000000')
r = p.add_run('Please approve and provide the controlled Google access below so the CardSnap backend can add submissions to the review Sheet and read Hala’s status updates. The application, security controls, staging deployment, named test accounts and automated tests are otherwise ready for the live integration check.')
set_font(r, size=9.2)

h = doc.add_paragraph(style='Heading 1')
h.add_run('Google access still required')
requirements = [
    'Create or copy the review Sheet in the approved Aventure Google Drive account for live use. The current Vision71 Sheet remains for testing only.',
    'Share that single Sheet as Editor with cardsnap@cardsnap-510210.iam.gserviceaccount.com. Access should not be granted to the wider Drive.',
    'Enable the Google Sheets API in Google Cloud project cardsnap-510210.',
    'Create a JSON key for the CardSnap service account and place its private_key in the staging and later approved production environment as GOOGLE_SERVICE_ACCOUNT_PRIVATE_KEY.',
    'Do not send the JSON key through chat, WhatsApp, email, browser code or repository files. Haroon or Hassan should enter it directly in the Vercel environment settings.',
    'Replace the Google Vision API key that was shared in chat, restrict the replacement to Cloud Vision only and store it as GOOGLE_VISION_API_KEY. CardSnap sends it in a request header, never in a URL.'
]
for item in requirements:
    add_bullet(doc, item)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(5)
r = p.add_run('Why this is needed. ')
set_font(r, size=8.8, bold=True, color='000000')
r = p.add_run('The Sheet URL identifies the spreadsheet but does not authorize private reads or writes. The service account email grants access and its private key proves the backend identity. A Google API key or OAuth client ID cannot replace this authentication for the private Sheet workflow.')
set_font(r, size=8.8)

h = doc.add_paragraph(style='Heading 1')
h.add_run('Version 1 production architecture recommendation')

data = [
    ('Login and authentication', 'CardSnap named accounts with scrypt password hashing, server sessions, Secure SameSite Strict cookies, CSRF protection, rate limits, lockout and immediate access removal.'),
    ('Database', 'MongoDB Atlas Free for the controlled pilot database. Tenant separation and AES 256 GCM encryption protect contact fields, OCR text and images.'),
    ('Image storage', 'Encrypted MongoDB cardImages storage for Version 1. Browser compression keeps images below 500 KB. Do not create Google Cloud Storage until separately approved.'),
    ('Cloud provider and region', 'Separate Vercel staging project now. Propose Mumbai for live application and database services when available, subject to Ali and Aventure approving hosting terms, region and data residency.'),
    ('Retention and deletion', 'Default image retention is 24 hours. Allowed range is 0 to 168 hours. The hourly application sweep is audited; MongoDB TTL is a two hour safety net.'),
    ('Backup method', 'Daily GitHub Actions export, compression and AES 256 GCM encryption with a separate backup key. Images, sessions and tokens are excluded. Encrypted artifacts are retained for 30 days. Restore testing has passed.'),
    ('Expected monthly cost', 'Expected internal test cost is 0 USD inside approved free allowances. Vercel Hobby is non commercial, and Google billing requirements must be confirmed before a live commercial pilot.'),
    ('Who can access data', 'Capturers see their own records. Hala reviews the Aventure owned Sheet. Vision71 Administrators manage accounts, retention, deletion and Approved CSV export. Support sees aggregate counts only and expires after 24 hours. No automatic Constant Contact connection exists in Version 1.')
]

table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(1.55)
table.columns[1].width = Inches(5.45)
set_table_borders(table)
hdr = table.rows[0].cells
hdr[0].width = Inches(1.55); hdr[1].width = Inches(5.45)
for i, text in enumerate(('Area', 'Recommendation')):
    hdr[i].text = ''
    set_cell_shading(hdr[i], '315C85')
    set_cell_margins(hdr[i], 85, 105, 85, 105)
    hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = hdr[i].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(text), size=8.2, bold=True, color='FFFFFF')
set_repeat_table_header(table.rows[0])

for idx, (area, recommendation) in enumerate(data):
    cells = table.add_row().cells
    cells[0].width = Inches(1.55); cells[1].width = Inches(5.45)
    if idx % 2:
        set_cell_shading(cells[0], 'F4F7FA'); set_cell_shading(cells[1], 'F4F7FA')
    for cell in cells:
        set_cell_margins(cell, 70, 105, 70, 105)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cells[0].text = ''; cells[1].text = ''
    p0 = cells[0].paragraphs[0]; p0.paragraph_format.space_after = Pt(0); p0.paragraph_format.line_spacing = 0.98
    set_font(p0.add_run(area), size=7.7, bold=True, color='000000')
    p1 = cells[1].paragraphs[0]; p1.paragraph_format.space_after = Pt(0); p1.paragraph_format.line_spacing = 0.98
    set_font(p1.add_run(recommendation), size=7.55)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(0)
r = p.add_run('Current status. ')
set_font(r, size=8.4, bold=True, color='000000')
r = p.add_run('The haroon branch is deployed to the private staging project. All 112 automated tests and the production build pass. The remaining live checks are Google Vision, Google Sheet write and status sync, and physical mobile testing after the credentials above are configured.')
set_font(r, size=8.4)

# Minimal footer
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.space_before = Pt(0)
set_font(fp.add_run('CardSnap by Vision71  |  30 September 2026'), size=7.2, color='6B7280')

# Metadata
props = doc.core_properties
props.title = 'CardSnap Google Sheets Access and Production Architecture'
props.subject = 'One page decision note for Ali Bhai'
props.author = 'Vision71'
props.keywords = 'CardSnap, Google Sheets, production architecture'

doc.save(OUT)
print(OUT)
