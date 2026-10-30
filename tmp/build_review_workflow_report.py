from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path(r'C:\Users\Administrator\Desktop\Aventure\CardLens\output\documents\Lead71_Post_Scan_Review_Workflow_Implementation_Report.docx')
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.58); sec.bottom_margin=Inches(.58); sec.left_margin=Inches(.68); sec.right_margin=Inches(.68)

def font(run,size=9.5,bold=False,italic=False):
    run.font.name='Arial'; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Arial'); run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial'); run.font.size=Pt(size); run.font.bold=bold; run.font.italic=italic; run.font.color.rgb=RGBColor(0,0,0); return run
for name,size,bold in [('Normal',9.5,False),('Title',20,True),('Heading 1',13,True),('Heading 2',10.5,True)]:
    st=doc.styles[name]; st.font.name='Arial'; st._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.space_after=Pt(4 if name=='Normal' else 5); st.paragraph_format.keep_with_next=name.startswith('Heading')
tp=doc.styles['Title']._element.get_or_add_pPr(); b=tp.find(qn('w:pBdr')); tp.remove(b) if b is not None else None

def para(text='',size=9.5,bold=False,italic=False,after=4):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.03; font(p.add_run(text),size,bold,italic); return p
def head(text,level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.space_before=Pt(6 if level==1 else 3); p.paragraph_format.space_after=Pt(3); font(p.add_run(text),13 if level==1 else 10.5,True); return p
def shade(cell,fill='E9E9E9'):
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),fill); cell._tc.get_or_add_tcPr().append(sh)
def celltext(cell,text,bold=False,size=8.2):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0; font(p.add_run(str(text)),size,bold); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers,rows,widths,size=8.2):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=False
    for i,h in enumerate(headers): celltext(t.rows[0].cells[i],h,True,size); shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,x in enumerate(row): celltext(c[i],x,False,size)
    for row in t.rows:
        for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def screenshot_box(label,height=1.0):
    t=doc.add_table(rows=1,cols=1); t.style='Table Grid'; t.autofit=False; t.columns[0].width=Inches(7.1); t.rows[0].height=Inches(height); cell=t.cell(0,0); celltext(cell,f'[Insert staging screenshot: {label}]\n\nBuild / date: ______________________________    Test case: ______________________________',False,9)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
def hyperlink(p,text,url):
    rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid); r=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'000000'); pr.append(color); u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); pr.append(u); r.append(pr); tx=OxmlElement('w:t'); tx.text=text; r.append(tx); h.append(r); p._p.append(h)

p=doc.add_paragraph(style='Title'); p.paragraph_format.space_after=Pt(2); font(p.add_run('Lead71 Post-Scan Review Workflow'),20,True)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(7); font(p.add_run('Implementation summary and staging evidence | 30 October 2026'),10)
para('To: Ali Bhai\nFrom: Ibrahim',after=5)
para('AoA Ali Bhai, the two-level post-scan review form has been implemented in the current Lead71 source for staging review. Fast Review now includes the essential contact and follow-up fields, while Add Sales Details starts collapsed. Follow-up action and due date appear only when Yes is selected. This report records what the source currently implements and leaves space for staging screenshots and measured results. The source review does not by itself confirm the deployed staging build or prove the three user tests; those evidence fields remain open below. Production has not been changed.',after=6)

head('Implementation summary')
table(['Blueprint requirement','Current implementation','Evidence status'],[
 ['Fast Review: Company Type','Visible and required; currently a text input with “Enter company type or Not sure” placeholder, not a dropdown.','Implemented in source; staging screenshot pending'],
 ['Fast Review: Contact Type','Visible and required dropdown: Prospect, Customer, Supplier, Partner, Other, Not sure.','Implemented in source; staging screenshot pending'],
 ['Follow-up Required','Visible required Yes / No radio choice.','Implemented in source; staging screenshot pending'],
 ['Next Action and Follow-up Due Date','Shown only after Yes; both become required in that case. Hidden and cleared from saved data when No.','Implemented in source; conditional staging proof pending'],
 ['Add Sales Details','Collapsed by default and expandable. Includes Product Interest, Aircraft Platform, Aviation Sector, Lead Priority, and Additional Meeting Notes.','Implemented in source; optional-save staging proof pending'],
 ['Save Contact','Fixed action in the modal footer. Basic required checks include event/source, a name or company, a contact method, Company Type, Contact Type, and Follow-up Required.','Implemented in source; staging save test pending'],
 ['Marketing permission','The form states saving does not subscribe a contact to marketing.','Visible in source; staging screenshot pending'],
], [1.55,3.75,1.8], 7.8)
para('Implementation note: the blueprint proposed Company Type as a quick-selection dropdown and suggested dropdown controls for several optional sales fields. The current form uses a required text input for Company Type; the optional sales fields are text inputs. These are differences from the blueprint and should be confirmed with Aventure before treating the field controls as final.',italic=True,after=5)

head('Fast Review field order and save rules')
para('The review panel shows Exhibition / Source, Contact Name, Company Name, Job Title, Email Address, Phone Number, Company Type, Contact Type, Follow-up Required, conditional Next Action and Follow-up Due Date, and optional Short Notes. Add Sales Details follows the Fast Review fields and is collapsed on first open. Save Contact remains in the modal footer.')
para('When Follow-up Required is Yes, the form requires both a Next Action and a Follow-up Due Date. When it is No, those fields are hidden and the submitted values are blanked. The optional Product Interest, Aircraft Platform, Aviation Sector, Lead Priority, and Additional Meeting Notes fields are not required by the form schema. Code review supports that they can be skipped; successful staging saves with them blank have not yet been evidenced here.')
para('The exhibition/source is initialized from OCR meeting context when present, otherwise from the first available exhibition option. We have not verified in staging whether event city/country, scan timestamp, and representative are visibly or reliably populated for this workflow. The form uses the existing Lead71 review modal and styling. Actual mobile scrolling and tap burden still need observation.')

head('Screenshot evidence')
para('Add screenshots from the staging build using fabricated contact data only. Include the default state, Follow-up Required = Yes, and the expanded optional section. Do not include real business-card images or personal contact information.',after=3)
screenshot_box('Fast Review default state with Add Sales Details collapsed')
screenshot_box('Follow-up Required = Yes showing Next Action and Due Date')
screenshot_box('Add Sales Details expanded and optional')

head('Three staging test cases')
para('The approved research PDF describes the original tap counts and durations as planning estimates. The table below is reserved for actual staging measurements after the review screen opens. No estimated value is presented as measured evidence.')
table(['Test case','Expected check','Actual taps / time / corrections','Result and usability notes'],[
 ['1. Basic contact; Follow-up = No','Review a complete synthetic contact, choose Company Type and Contact Type, select No, save with optional sales fields blank.','Taps: ______  Time: ______\nCorrections: ______','Save: ______\nIssues: __________________'],
 ['2. Qualified lead; Follow-up = Yes','Select Yes; confirm both conditional fields appear and are required; enter action/date; optionally add details; save.','Taps: ______  Time: ______\nCorrections: ______','Save: ______\nIssues: __________________'],
 ['3. Incomplete / uncertain contact','Correct only clear OCR errors; use Not sure where needed; skip optional fields; save without inventing missing information.','Taps: ______  Time: ______\nCorrections: ______','Save: ______\nIssues: __________________'],
], [1.35,2.55,1.55,1.65], 7.4)
para('Measurement method: start timing when Fast Review is visible and stop after the staging save confirmation. Count user taps after the screen opens. Count each manually corrected field once. Record device, viewport, build/commit, date, and whether the user completed the task without assistance. Repeat each case and record every run rather than averaging away usability problems.',italic=True,after=4)
screenshot_box('Staging save confirmation using synthetic data')

head('Validation status and remaining points')
table(['Item','Current status'],[
 ['Two-level form structure and field rendering','Present in local source; confirm deployed staging revision with commit/PR and screenshots.'],
 ['Conditional follow-up controls and validation','Present in source; verify Yes/No behavior and saved values in staging.'],
 ['Optional fields can be skipped','Form schema marks sales-detail fields optional; successful staging save with blanks remains to be recorded.'],
 ['Three user test cases','Not recorded in this report. Complete the measurements above using synthetic contacts.'],
 ['Tap counts, completion time, corrections, usability','Not yet measured; enter actual observations from staging, not the original PDF estimates.'],
 ['Blueprint field-control differences','Company Type and optional sales controls differ from the proposed dropdown controls; confirm whether to retain or revise them.'],
 ['Production','No production change is included or approved by this report. Review staging evidence before any production decision.'],
], [2.15,4.95], 8)

head('Research sources')
para('Official product documentation reviewed for the approved research blueprint. Accessed 30 October 2026.',after=3)
for label,url in [
 ('Cvent LeadCapture','https://www.cvent.com/en/event-marketing-management/lead-capture'),
 ('HubSpot mobile contact creation and business-card scanning','https://knowledge.hubspot.com/import-and-export/import-contacts-on-the-hubspot-mobile-app'),
 ('HubSpot conditional property logic','https://knowledge.hubspot.com/properties/set-up-conditional-logic-for-enumeration-properties'),
 ('fielddrive Leads lead retrieval','https://www.fielddrive.com/solution/lead-retrieval'),
]:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.14); p.paragraph_format.space_after=Pt(2); font(p.add_run('• '+label+': '),8.7); hyperlink(p,url,url); font(p.add_run(' (accessed 30 October 2026)'),8.7)

head('Recommendation')
para('The two-level review workflow is implemented in the current source and is ready for staging evidence collection. Complete the three staging test cases, verify optional-field saves, capture the screenshots above, confirm the Company Type control with Aventure, and attach the staging commit reference. Do not change production until the staging evidence is reviewed and approved.',after=3)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(0); font(p.add_run('Staging review pending; production deployment not approved.'),12,True)

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(footer.add_run('Lead71 Post-Scan Review Workflow  |  '),8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title='Lead71 Post-Scan Review Workflow Implementation Report'; doc.core_properties.subject='Staging implementation and evidence'; doc.core_properties.author='Lead71'; doc.core_properties.keywords='Lead71, Fast Review, staging, follow-up'
OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)
