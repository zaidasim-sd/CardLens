from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = r"C:\Users\Administrator\Desktop\Aventure\CardLens\output\documents\Lead71_Gemini_AI_Mapping_Revised_Assessment_Final.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.62)
sec.bottom_margin = Inches(.62)
sec.left_margin = Inches(.72)
sec.right_margin = Inches(.72)

def font(run, size=10, bold=False, italic=False):
    run.font.name = 'Arial'
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Arial')
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Arial')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    return run

styles = doc.styles
for stylename, size, bold in [('Normal', 9.5, False), ('Title', 20, True), ('Heading 1', 14, True), ('Heading 2', 11, True)]:
    st = styles[stylename]
    st.font.name = 'Arial'
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(4 if stylename == 'Normal' else 6)
    st.paragraph_format.keep_with_next = stylename.startswith('Heading')

# Word's built-in Title style carries a blue bottom border in this Office build.
# Strip that inherited border so the requested report stays plain black.
title_style_ppr = styles['Title']._element.get_or_add_pPr()
title_bdr = title_style_ppr.find(qn('w:pBdr'))
if title_bdr is not None:
    title_style_ppr.remove(title_bdr)

def shade(cell, fill='F2F2F2'):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); pr.append(shd)

def cell_text(cell, text, bold=False, size=8.2):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    font(p.add_run(text), size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None, size=8.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    t.autofit = False
    for i,h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, True, size)
        shade(t.rows[0].cells[i], 'E7E7E7')
    for row in rows:
        cells = t.add_row().cells
        for i,v in enumerate(row): cell_text(cells[i], v, False, size)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def p(text='', boldlead=None, italic=False, after=4):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.05
    if boldlead and text.startswith(boldlead):
        font(para.add_run(boldlead), bold=True)
        font(para.add_run(text[len(boldlead):]), italic=italic)
    else: font(para.add_run(text), italic=italic)
    return para

def heading(text, level=1):
    para = doc.add_paragraph(style=f'Heading {level}')
    para.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    para.paragraph_format.space_after = Pt(3)
    font(para.add_run(text), 14 if level == 1 else 11, True)
    return para

def bullet(text):
    para=doc.add_paragraph(style='List Bullet')
    para.paragraph_format.left_indent=Inches(.22)
    para.paragraph_format.first_line_indent=Inches(-.14)
    para.paragraph_format.space_after=Pt(2)
    font(para.add_run(text), 9.2)
    return para

title=doc.add_paragraph(style='Title')
title.paragraph_format.space_after=Pt(3)
font(title.add_run('Lead71 Gemini AI Mapping Assessment'),20,True)
sub=doc.add_paragraph(); sub.paragraph_format.space_after=Pt(7)
font(sub.add_run('Staging evidence and production readiness | 30 October 2026'),10)
p('To: Ali Bhai\nFrom: Haroon and Ibrahim', after=6)
p('Ali Bhai, we have updated the assessment to reflect the current code and the staging status reported by the team. The manual-parser fallback is reported as deployed to staging. The staging deployment reference and screenshot still need to be attached, so this report does not treat that deployment as independently evidenced. No production API key, production deployment, or Aventure card data is approved by this assessment.', after=7)

heading('Current status')
table(['Area','Status','What the evidence supports'],[
 ['AI classification with manual fallback','Reported implemented in staging; staging evidence pending','Current source contains the AI-to-manual fallback path and a local automated test for upstream failure, missing key, and invalid AI output. The live staging commit and screenshot are not present here.'],
 ['Staging verification and synthetic test results','Pending test','No staging test record, processing-time measurements, or concurrency results were available for review.'],
 ['Application request caps','Implemented locally only; staging configuration pending','Code defaults are 120 requests per user per hour, 240 per IP per hour, and 900 OCR requests per tenant per month. The monthly cap covers OCR requests in all classification modes.'],
 ['Immediate AI rollback switch','Proposed but not built','The current source does not contain ENABLE_AI_MAPPING=false. No working-switch proof can be supplied.'],
 ['Production use of real Aventure card data','No-Go for production','Deployment, data-cleanup, cost-limit, rollback, and staging test evidence remain incomplete.']
], [1.42,1.65,4.0], 8)
p('Status labels in this report distinguish code reviewed locally from behavior observed in staging. “Reported in staging” records the team’s statement; it is not a substitute for a commit reference or test evidence.', italic=True, after=6)

heading('Fallback behavior')
p('In AI mode, the server sends the OCR text to Gemini and checks that the returned fields contain a readable contact. If classification fails, times out, returns an invalid result, or cannot be used because the key/configuration is unavailable, the current code runs the manual OCR parser on the same OCR text. A readable parsed contact is returned for review; the saved record is not discarded by this fallback path.')
p('On the review screen, the fallback result is labeled “Manual classification fallback.” The application does not show the Gemini error to the user in this flow. The current default AI timeout is 8 seconds. There is no automatic Gemini retry: the user can correct the review fields and continue to save, or rescan the card to start a new OCR and AI request. If OCR cannot read the card or neither parser produces a readable contact, the current user message is: “We could not read this card. Capture the card again or enter the details manually.”')
p('The current code therefore supports fallback without losing the OCR text during classification failure. A completed staging demonstration is still needed to prove the visible message, review form, and successful save end to end.', after=6)

heading('Staging screenshot evidence')
p('Insert a screenshot here showing the staging review screen after an AI classification failure. The image should show the “Manual classification fallback” label and the retained fields in the review form. Add a second screenshot or test record showing that the reviewer saved the corrected record successfully. Use synthetic test content only.', italic=True, after=2)
box=doc.add_table(rows=1, cols=1); box.style='Table Grid'; box.autofit=False; box.columns[0].width=Inches(7.05); box.rows[0].height=Inches(1.65)
cell_text(box.cell(0,0),'[Insert staging fallback screenshot here]\n\nTest ID / date: ______________________________    Staging URL / build: ______________________________',False,9)
doc.add_paragraph().paragraph_format.space_after=Pt(1)

heading('Required synthetic test record')
p('The revised source report’s three successful examples are not repeated as test evidence. Complete this table with synthetic card text or synthetic printed sample cards only. “Pending” means the run has not been evidenced in this report; no timing or outcome has been invented.')
table(['Scenario','Measured end-to-end time','Mapping and correction','Fallback / save result'],[
 ['Gemini failed API response','Pending staging measurement','Pending test; record mapped fields and manual correction','Pending test; record fallback label and save outcome'],
 ['Gemini timeout (configured default 8 s)','Pending staging measurement','Pending test; record mapped fields and manual correction','Pending test; record fallback label and save outcome'],
 ['Invalid or incomplete AI response','Pending staging measurement','Locally unit-tested; synthetic staging result pending','Pending staging proof of fallback, correction, and save'],
 ['Gemini quota or rate-limit response','Pending staging measurement','Pending test; record mapped fields and manual correction','Pending test; upstream limit behavior not observed'],
 ['Poor or unreadable OCR input','Pending staging measurement','Pending test; record OCR text quality and correction needed','Pending test; record manual entry or retake and save path'],
 ['Concurrent scanning','Pending load-test measurement','Pending test; record requests, mappings, and errors','Pending test; record any fallback or rejected requests']
], [1.55,1.25,2.05,2.05], 7.6)
p('Attach sanitized run logs or a test sheet extract here. Do not include card images, names, email addresses, phone numbers, or other real contact details.', italic=True)
box=doc.add_table(rows=1, cols=1); box.style='Table Grid'; box.autofit=False; box.columns[0].width=Inches(7.05); box.rows[0].height=Inches(.9)
cell_text(box.cell(0,0),'[Add synthetic test results, timestamps, and staging evidence here]',False,9)
doc.add_paragraph().paragraph_format.space_after=Pt(1)
p('Local code evidence: the checked-in OCR handler tests cover successful AI mapping, upstream failure, missing key, invalid AI output, and comparison mode with synthetic text. These are automated local tests, not staging end-to-end measurements. No evidence for a real quota exhaustion, prepaid-balance exhaustion, poor-image OCR, or concurrent staging load run was available for this revision.', after=6)

heading('Cost and request limits')
p('The live billing position provided for this assessment is a $20 prepaid Gemini balance with auto-reload off. We did not inspect the billing console, so that amount and setting remain owner-reported. The earlier $50 “hard cap” statement is withdrawn: no such enforced application or provider cap was evidenced. A normal billing budget is an alerting control and is not described here as a hard stop.')
p('In the current source, the application enforces a default of 120 OCR requests per signed-in user per hour, 240 per source IP per hour, and 900 OCR requests per tenant per month. The user and IP limits can be overridden with server environment variables; the monthly limit is set with OCR_MONTHLY_CAP. The monthly cap applies to the OCR endpoint across AI and manual modes. When an app cap is exceeded, the handler rejects the request with HTTP 429 and returns the generic OCR message. Staging values have not been verified.')
p('When Gemini returns an error during AI classification, the handler attempts the manual-parser fallback. If it finds readable contact fields, the user can review and save them. Google’s published billing guide says a zero prepaid balance stops Gemini API requests and returns HTTP 402 until credits are added; the application’s behavior for that exact staging response has not been demonstrated. The same is true of actual Gemini quota/rate-limit responses. No staging test has confirmed prepaid exhaustion, quota exhaustion, or the app-cap response.', after=6)

heading('Privacy and data handling')
p('The code path sends card image bytes to Google Cloud Vision for OCR. The Gemini classification request contains the resulting OCR text; the Gemini request body is text-only and does not include the card image. This describes the current code, not a verified network capture from staging.')
p('Google’s published Gemini API terms say that when Gemini API access is through a Cloud project with an active billing account, prompts and responses are not used to improve Google products. Google also states that prompts and responses may be logged for abuse monitoring for up to 55 days, and that flagged content may be reviewed by authorized personnel. The owner-reported prepaid balance suggests paid access, but the exact project’s active billing status and applicable terms have not been checked in the console. This assessment does not claim an enterprise agreement.')
p('Suggested client disclosure, subject to confirming the account and reviewing applicable contract terms: “Lead71 uses Google Cloud Vision to extract text from business-card images. If AI classification is enabled, extracted text, rather than the card image, is sent to the Gemini API to organize contact fields. For paid Gemini API use through a project with active Cloud Billing, Google states prompts and responses are not used to improve its products. Google may retain prompts and responses for abuse monitoring for up to 55 days.”')
p('The source code emits a structured fallback event with a reason code and tenant identifier; it is designed not to include card text or upstream response bodies. That does not establish that all staging logs are free of personal data. Log contents still need a staging review.', after=6)

heading('Deployment and cleanup evidence still required')
table(['Evidence item','Current finding','Reference / attachment'],[
 ['Staging commit or PR','Not provided; team reports fallback is deployed','____________________________________________'],
 ['Separate staging and production environments / databases','Not confirmed for this AI change','____________________________________________'],
 ['Gemini API key stored only as a server-side staging secret','The source expects GEMINI_API_KEY server-side; staging secret scope not verified','____________________________________________'],
 ['ENABLE_AI_MAPPING=false rollback test','Switch is absent from current source; not tested','Implement, deploy, and attach proof before production consideration'],
 ['Prior report sharing status','Not verifiable from this workspace','Confirm recipients and whether any copy was forwarded'],
 ['Staging records, test Sheets, and logs cleanup','Not verifiable from this workspace','Confirm deletion and attach dated evidence; check logs separately']
], [1.75,2.7,2.45], 7.8)
p('The prior report contained visible card imagery and personal/contact details. This revised report omits the images and those details. The prior copy’s distribution and any records derived from it have not been verified here; Haroon and Zaid should record the distribution check and cleanup evidence above. Do not run cleanup against production data.', after=6)

heading('Actions before reconsidering production')
for item in [
 'Attach the staging commit or PR and confirm the staging database, Sheets, and secrets are isolated from production.',
 'Run each synthetic scenario in the test table. Record request time, result mapping, required correction, fallback outcome, and whether saving succeeded.',
 'Demonstrate zero-balance, quota/rate-limit, and application request-cap behavior. Keep the $20 prepaid balance and auto-reload status documented from the billing console.',
 'Implement and test an actual server-side AI disable switch, then attach evidence that setting ENABLE_AI_MAPPING=false bypasses Gemini immediately.',
 'Check whether the earlier report was forwarded and document cleanup of related staging records, test Sheets, and logs.',
 'Confirm the exact billing project and applicable Gemini terms before issuing a client privacy disclosure.'
]: bullet(item)

heading('Recommendation')
p('No production API key, Vercel production deployment, or Aventure card data should be used until the evidence above is complete and reviewed. The current record does not prove production safeguards or end-to-end staging behavior.')
rec=doc.add_paragraph(); rec.paragraph_format.space_before=Pt(3); rec.paragraph_format.space_after=Pt(0)
font(rec.add_run('No-Go for production'),14,True)

# Page number field in a quiet footer.
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
font(footer.add_run('Lead71 Gemini AI Mapping Assessment  |  '),8)
r=footer.add_run(); font(r,8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)

doc.core_properties.title='Lead71 Gemini AI Mapping Assessment'
doc.core_properties.subject='Staging evidence and production readiness'
doc.core_properties.author='Lead71'
doc.core_properties.keywords='Gemini, OCR, staging, evidence'
import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(OUT)
